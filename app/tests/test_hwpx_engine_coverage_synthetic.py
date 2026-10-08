"""Engine cell grants: positive coverage and destructive counterexamples."""
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
import pytest
from docx import Document
from PIL import Image
from core.docx.services.hwpx_analysis_adapter import index_hwpx_structure
from core.docx.services.hwpx_protected_regions import (
    authorize_guidance_value_writes,find_record_row_targets,is_value_cell_guidance,
    assess_fields,document_parts,normalized_cell_label,
)
from core.docx.services.hwpx_fill import commit_guidance_value_writes,fill_hwpx,_conditional_guidance_value
from auto_write.services.hwpx_narrative_source import direct_cell_plan,narrative_from_references
from auto_write.services.hwpx_form_diff import compare_hwpx_forms
from auto_write.services.hwpx_pic_insert import insert_cell_reference_images
from test_hwpx_round2_fill import _cell,_run,_row,_sec,_write
HP='http://www.hancom.co.kr/hwpml/2011/paragraph'
NS={'hp':HP}
LABELS=['아이디어 제안배경','신청자 역량','창업 아이디어','핵심내용','우수성 및 차별성','시장 경쟁력','기대효과']

def table(rows):
    return '<hp:p><hp:run><hp:tbl>'+''.join(rows)+'</hp:tbl></hp:run></hp:p>'

def title(text):
    return table(['<hp:tr>'+_cell(0,0,_run(text))+'</hp:tr>'])

def package(tmp_path,body):
    p=tmp_path/'source.hwpx';_write(p,_sec(body));return p

def root(path):
    with ZipFile(path) as z:return etree.fromstring(z.read('Contents/section0.xml'))

def references(path,images=2,plain=False):
    doc=Document()
    for k,v in {'대표자':'홍길동','팀명':'테스트팀','팀원 수':'4명(대표 포함)','구분':'예비창업자','사업자등록번호':'해당없음','아이디어명':'테스트 시스템','지식재산권':'특허(출원) / 테스트 시스템 / 10-2026-0000000 / 홍길동 / 2026.01.01 출원'}.items():
        doc.add_paragraph(f'{k}: {v}')
    for i,label in enumerate(LABELS):
        if plain:doc.add_paragraph(label)
        else:doc.add_heading(label,level=1)
        for n in range(3):doc.add_paragraph(f'{label} 합성 본문 {n+1}: 제공된 설명으로 확인한다.')
        if i<images:
            blob=BytesIO();Image.new('RGB',(600,300),'blue').save(blob,format='PNG');blob.seek(0)
            doc.add_paragraph().add_run().add_picture(blob)
    doc.save(path);return path

def narrative_table():
    rows=[]
    for i,label in enumerate(LABELS):
        value=f'{label}에 대해 작성' if i in {0,1,2,6} else ''
        row=_row(i,label,value).replace('<hp:subList>','<hp:cellSz width="10000" height="1000"/><hp:subList>')
        rows.append(row)
    return table(rows)

def apply(tmp_path,src,identity,refs=(),meta=None):
    index=index_hwpx_structure(src)
    grants,values,notes,images=direct_cell_plan(index,identity,refs,meta or {})
    output=tmp_path/'filled.hwpx'
    written=commit_guidance_value_writes(src,output,grants,values)
    return output,written,notes,images

@pytest.mark.parametrize('value',['(문서 자동 작성 서비스)','올해 사업계획서를 작성','(사업계획서 작성 완료)','2026.01.01','※ 작성 완료: 이미 제출함','(테스트팀)','기재 완료','서비스를 소개'])
def test_actual_values_never_instruction(value):
    assert not is_value_cell_guidance(value)

@pytest.mark.parametrize('value',['본인이 희망하는 내용을 기재','핵심내용에 대해 작성','기대되는 영향력 등 기재','등록된 사업자가 있는 경우 기재','팀으로 참가시에만 기재'])
def test_qualified_instruction(value):
    if value=='본인이 희망하는 내용을 기재':
        assert is_value_cell_guidance(value)
    else:assert is_value_cell_guidance(value)

@pytest.mark.parametrize('identity,expected',[({'구분':'예비창업자','기업명':'잘못 쓰면 안됨','팀명':'테스트팀','팀원 수':'4명'},['해당없음']*3+['테스트팀','4명']),({'사업자등록번호':'123-45-67890','사업자명':'등록 상호','팀명':'테스트팀','팀원수':'4명'},['등록 상호','123-45-67890','','테스트팀','4명']),({'기업명':'쓰기 금지'},['등록된 사업자가 있는 경우 기재']*3+['팀으로 참가시에만 기재']*2)])
def test_conditional_fields(tmp_path,identity,expected):
    labels=['사업자명','사업자등록번호','사업개시일','팀명','팀원 수']
    src=package(tmp_path,table([_row(i,l,'등록된 사업자가 있는 경우 기재' if i<3 else '팀으로 참가시에만 기재') for i,l in enumerate(labels)]))
    original=src.read_bytes();out,written,notes,_=apply(tmp_path,src,identity)
    actual=[''.join(c.itertext()) for c in root(out).xpath('//hp:tc[position()=2]',namespaces=NS)]
    for i,e in enumerate(expected):
        if e:assert e==actual[i]
    assert '잘못 쓰면 안됨' not in ''.join(root(out).itertext())
    assert src.read_bytes()==original
    if not any(k in identity for k in ('구분','사업자등록번호')):assert any('UNDECIDED' in n for n in notes)


def test_legacy_conditional_never_company_alias(tmp_path):
    src=package(tmp_path,table([_row(0,'사업자명','등록된 사업자가 있는 경우 기재')]))
    out=tmp_path/'legacy.hwpx';fill_hwpx(src,out,identity={'기업명':'가짜상호','사업자등록번호':'해당없음'})
    assert '가짜상호' not in ''.join(root(out).itertext())


def test_parts_fanout_and_same_part_duplicate(tmp_path):
    body=''.join(title(t)+table([_row(0,'아이디어명',''),_row(1,'신청자명(팀명)','')]) for t in ('신청서','사업계획서','참가 서약서'))
    src=package(tmp_path,body)
    index=index_hwpx_structure(src);assert len(set(document_parts(index).values()))==3
    out,_,_,_=apply(tmp_path,src,{'아이디어명':'테스트 시스템','대표자':'홍길동','팀명':'테스트팀'})
    text=''.join(root(out).itertext());assert text.count('테스트 시스템')==3;assert text.count('홍길동 (팀명: 테스트팀)')==3
    src=package(tmp_path,title('신청서')+table([_row(0,'팩스',''),_row(1,'팩스','')]))
    assert not authorize_guidance_value_writes(index_hwpx_structure(src))


def record_table():
    headers=['권리구분','등록(출원)명칭','등록(출원)번호','등록(출원)인','기타']
    label=_cell(0,0,_run('지식재산권')).replace('rowSpan="1"','rowSpan="2"')
    return table(['<hp:tr>'+label+''.join(_cell(i+1,0,_run(h)) for i,h in enumerate(headers))+'</hp:tr>','<hp:tr>'+''.join(_cell(i+1,1,_run('')) for i in range(5))+'</hp:tr>']),headers


def test_explicit_record_row_and_overflow(tmp_path):
    body,headers=record_table();src=package(tmp_path,body)
    assert len(find_record_row_targets(index_hwpx_structure(src)))==5
    values=['특허','테스트 시스템','10-2026-0000000','홍길동','출원']
    meta={'hwpx_records':{'지식재산권':[dict(zip(headers,values))]*2}}
    out,written,notes,_=apply(tmp_path,src,{},meta=meta)
    assert all(v in ''.join(root(out).itertext()) for v in values);assert len(written)==5
    assert any('OVERFLOW' in n for n in notes)
    out,written,notes,_=apply(tmp_path,src,{'지식재산권':'a / b'})
    assert written=={};assert any('FIELD_COUNT_MISMATCH' in n for n in notes)


def test_seven_narratives_scope_style_and_provided_images(tmp_path):
    src=package(tmp_path,narrative_table());original=src.read_bytes();ref=references(tmp_path/'ref.docx')
    out,written,notes,images=apply(tmp_path,src,{},[ref])
    assert set(written)==set(LABELS);assert len(images)==2
    old,new=root(src),root(out)
    for a,b in zip(old.xpath('//hp:tc[position()=1]',namespaces=NS),new.xpath('//hp:tc[position()=1]',namespaces=NS)):
        assert etree.tostring(a)==etree.tostring(b)
    assert '에 대해 작성' not in ''.join(new.itertext())
    assert not new.xpath('//hp:linesegarray',namespaces=NS)
    with ZipFile(src) as z,ZipFile(out) as o:
        for member in z.namelist():
            if member not in {'Contents/section0.xml','Contents/header.xml'}:assert z.read(member)==o.read(member)
        old_head=etree.fromstring(z.read('Contents/header.xml'));new_head=etree.fromstring(o.read('Contents/header.xml'))
        assert all(etree.tostring(p)==etree.tostring(q) for p,q in zip(old_head[0][0],new_head[0][0]))
    assert src.read_bytes()==original
    final=tmp_path/'images.hwpx';insert_cell_reference_images(out,final,images,authorization_source=src,values=written)
    pics=root(final).xpath('//hp:pic',namespaces=NS);assert len(pics)==2
    image_paragraphs = [pic.getparent().getparent() for pic in pics]
    image_ids = [p.get('id') for p in image_paragraphs]
    assert len(set(image_ids)) == len(image_ids)
    for paragraph_id in image_ids:
        assert len(root(final).xpath('//hp:p[@id=$pid]', namespaces=NS, pid=paragraph_id)) == 1
    for pic in pics:
        assert pic.find(f'{{{HP}}}pos').get('treatAsChar')=='1'
        size=pic.find(f'{{{HP}}}sz');assert 0<int(size.get('width'))<=10000;assert int(size.get('height'))>0
    with ZipFile(final) as z:
        assert len([n for n in z.namelist() if n.startswith('BinData/')])==2
        manifest=etree.fromstring(z.read('Contents/content.hpf'))
        assert len(manifest.xpath('//*[local-name()="item" and @isEmbeded="1"]'))==2
    rep=compare_hwpx_forms(src,final)
    assert rep.form_intact,rep.as_dict();assert rep.guidance_replacements==4;assert rep.value_paragraphs>=14


def test_plain_heading_exact_matching_and_missing_source(tmp_path):
    src=package(tmp_path,narrative_table());ref=references(tmp_path/'plain.docx',images=0,plain=True)
    out,written,notes,_=apply(tmp_path,src,{},[ref]);assert set(written)==set(LABELS)
    out,written,notes,_=apply(tmp_path,src,{});assert not written;assert len([n for n in notes if 'NO_SOURCE' in n])==7


def test_stale_or_forged_grant_rejected(tmp_path):
    src=package(tmp_path,narrative_table());grant=authorize_guidance_value_writes(index_hwpx_structure(src))[0]
    out=tmp_path/'refused.hwpx';out.write_bytes(b'keep')
    with pytest.raises(ValueError,match='STALE'):
        commit_guidance_value_writes(src,out,[replace(grant,expected_raw_text='wrong')],{grant.field_label:'본문'})
    assert out.read_bytes()==b'keep'
    with pytest.raises(ValueError,match='덮어쓰기'):commit_guidance_value_writes(src,src,[grant],{grant:'본문'})


def test_choice_marks_only_and_label_geometry_destruction(tmp_path):
    src=package(tmp_path,table([_row(0,'구분','□ 예비창업자'),_row(1,'분야','□ 기술')]))
    def changed(path,change):
        with ZipFile(src) as a,ZipFile(path,'w') as b:
            for n in a.namelist():
                data=a.read(n)
                if n=='Contents/section0.xml':data=change(data.decode()).encode()
                b.writestr(n,data)
        return compare_hwpx_forms(src,path)
    rep=changed(tmp_path/'choice.hwpx',lambda s:s.replace('□','■'))
    assert rep.form_intact and rep.choice_marks==2
    assert not changed(tmp_path/'label.hwpx',lambda s:s.replace('구분','변경라벨')).form_intact
    assert not changed(tmp_path/'geometry.hwpx',lambda s:s.replace('colAddr="1"','colAddr="2"')).form_intact
    assert not changed(tmp_path/'table.hwpx',lambda s:s.replace('hp:tbl','hp:removed')).form_intact


@pytest.mark.parametrize('failures,generated',[(1,True),(2,False)])
def test_pdf_exactly_one_timeout_retry(monkeypatch,tmp_path,failures,generated):
    from core.docx.services import submission_gates as gate,hwp_docx_convert as com
    src=tmp_path/'form.hwpx';src.write_bytes(b'source');dest=tmp_path/'form.pdf'
    calls=[];sleeps=[]
    def export(a,b):
        calls.append((a,b))
        if len(calls)<=failures:raise com.HangulComTimeout('Dispatch',30)
        b.write_bytes(b'%PDF-test')
    monkeypatch.setattr(gate.sys,'platform','win32');monkeypatch.setattr(com,'hancom_com_available',lambda:True)
    monkeypatch.setattr(com,'export_pdf_via_com',export)
    import time
    monkeypatch.setattr(time,'sleep',lambda seconds:sleeps.append(seconds))
    result=gate._try_hangul_com_pdf(src,dest)
    assert result.generated==generated and result.attempts==2 and len(calls)==2
    assert sleeps==[5] and len(result.attempt_reasons)==2
    assert src.read_bytes()==b'source'


@pytest.mark.parametrize('configured,expected',[('10',10),('180',180),('45',45),('9',30),('181',30),('nan',30),('invalid',30)])
def test_dispatch_config_only(monkeypatch,configured,expected):
    from core.docx.services.hwp_docx_convert import _com_stage_timeout
    monkeypatch.setenv('AUTO_WRITE_HANGUL_DISPATCH_TIMEOUT',configured)
    assert _com_stage_timeout('Dispatch')==expected
    assert _com_stage_timeout('SaveAs[PDF]')==120


def test_empty_narrative_title_and_no_form_objects_masked(tmp_path):
    body=table(['<hp:tr>'+_cell(0,0,_run('사업계획서')).replace('colSpan="1"','colSpan="2"')+'</hp:tr>',_row(1,'핵심내용','')])
    src=package(tmp_path,body);ref=references(tmp_path/'plain.docx',images=0)
    out,written,_,_=apply(tmp_path,src,{},[ref]);assert '핵심내용 합성 본문' in written['핵심내용']
    controlled=tmp_path/'control.hwpx'
    _write(controlled,_sec(table([_row(0,'핵심내용','').replace('<hp:t></hp:t>','<hp:checkBtn name="old"/><hp:t></hp:t>',1)])))
    with ZipFile(controlled) as a,ZipFile(out,'w') as b:
        for n in a.namelist():
            data=a.read(n)
            if n=='Contents/section0.xml':data=data.replace(b'name="old"',b'name="changed"').replace(b'<hp:t></hp:t>', '<hp:t>본문</hp:t>'.encode())
            b.writestr(n,data)
    assert not compare_hwpx_forms(controlled,out).form_intact


def test_image_requires_reissued_origin_and_exact_current_body(tmp_path):
    src=package(tmp_path,narrative_table());ref=references(tmp_path/'ref.docx',images=1)
    out,written,_,images=apply(tmp_path,src,{},[ref]);final=tmp_path/'refused.hwpx'
    with pytest.raises(ValueError,match='AUTHORIZATION_REQUIRED'):insert_cell_reference_images(out,final,images)
    forged={replace(g,source_sha256='bad'):v for g,v in images.items()}
    with pytest.raises(ValueError,match='STALE'):insert_cell_reference_images(out,final,forged,authorization_source=src,values=written)
    with pytest.raises(ValueError,match='TEXT_CHANGED'):insert_cell_reference_images(out,final,images,authorization_source=src,values={k:'wrong' for k in written})
    assert not final.exists()
