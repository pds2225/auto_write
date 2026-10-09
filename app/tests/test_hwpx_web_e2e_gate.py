# -*- coding: utf-8 -*-
"""Gate: 웹/API HWPX 채움·제출이 #198–#205 계약을 유지하는지.

생산 경로는 FastAPI ``auto_write.main:app`` 이다. HWPX 업로드는 DOCX로
바꾸지 않고, 생성은 ``ProjectService._generate_hwpx_direct`` →
``submit_hwpx(preserve_template=True)`` 이다.

이 표면:
- POST /api/templates
- POST /api/projects
- POST /projects/{project_id}/generate
- POST /api/projects/{project_id}/generate
- GET /api/projects/{project_id}/artifacts
- GET /downloads/{project_id}/{artifact_name}
- POST /console/documents/write
- GET /console/results/{project_id}
- GET /console/results/{project_id}/download/{name}

운영 콘솔의 cross_form 전사는 DOCX 타깃만 한다. HWPX 타깃은
``hwpx_direct_fill`` 로 같은 direct-fill을 탄다.
"""

from __future__ import annotations

import hashlib
import zipfile
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse

import pytest
from docx import Document
from fastapi.testclient import TestClient
import auto_write.main as main
import auto_write.operator_main  # noqa: F401  — 콘솔 라우트를 같은 app에 붙인다
from auto_write.services.hwp_docx_convert import hancom_com_available
from auto_write.services.hwpx_acceptance import run_hwpx_acceptance
from auto_write.services.hwpx_submit import (
    RHWP_DISABLED_RENDER_NOTE,
    RHWP_DISABLED_REPAIR_NOTE,
)
from core.docx.services.hwpx_analysis_adapter import index_hwpx_structure
from core.docx.services import native_hwp
from core.docx.services.hwpx_protected_regions import assess_fields
from test_hwpx_cross_feature_integration import (
    _all_texts,
    _assert_structure_held,
    _at,
    _body_runs,
    _cells,
    _foreign_t,
    _hwpx as _write_combined,
    _identity,
    _member,
)
from test_hwpx_hancom_reopen_gate import (
    _CHECKBOX,
    _DATE,
    _EXISTING,
    _FILL,
    _SIGNATURE,
    _assert_package,
    _force_rhwp_absent,
    _section_clean,
    _section_protected_broken,
    _write,
)

_HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"

_HWPX_WEB_ROUTES = (
    ("POST", "/api/templates"),
    ("POST", "/api/projects"),
    ("POST", "/projects/{project_id}/generate"),
    ("POST", "/api/projects/{project_id}/generate"),
    ("GET", "/api/projects/{project_id}/artifacts"),
    ("GET", "/downloads/{project_id}/{artifact_name}"),
    ("POST", "/console/documents/write"),
    ("GET", "/console/results/{project_id}"),
    ("GET", "/console/results/{project_id}/download/{name}"),
)


@pytest.fixture
def web(monkeypatch, tmp_path):
    _force_rhwp_absent(monkeypatch)

    def _complete_native_review(path: str) -> dict:
        candidate = Path(path)
        digest = hashlib.sha256(candidate.read_bytes()).hexdigest()
        return {
            "ok": True,
            "severity": "PASS",
            "renderer": "test-evidence",
            "render_status": "PASS",
            "reopen_status": "PASS",
            "visual_review": "PASS",
            "candidate_sha256": digest,
            "render_source_sha256": digest,
            "page_count": 1,
            "pixel_reopen_claimed": False,
            "l005_pixel": "JUDGMENT",
            "l050_pdf": "NOT_SIBLING",
            "message": "synthetic web E2E complete render evidence",
        }

    monkeypatch.setattr(native_hwp, "verify_hwpx_native", _complete_native_review)
    root = tmp_path / "workspace"
    settings = replace(
        main.storage.settings,
        workspace_root=root,
        template_root=root / "templates",
        project_root=root / "projects",
        results_root=tmp_path / "results",
    )
    for path in (settings.template_root, settings.project_root, settings.results_root):
        path.mkdir(parents=True)
    monkeypatch.setattr(main.storage, "settings", settings)
    return main


@pytest.fixture
def client(web):
    return TestClient(web.app, raise_server_exceptions=False)


def test_hwpx_web_routes_are_the_fill_submit_surface():
    found = set()
    for route in main.app.routes:
        methods = getattr(route, "methods", None) or set()
        path = getattr(route, "path", None)
        if not path:
            continue
        for method in methods:
            found.add((method, path))
    missing = [item for item in _HWPX_WEB_ROUTES if item not in found]
    assert missing == []


def _loc(response) -> str:
    parsed = urlparse(response.headers["location"])
    query = f"?{parsed.query}" if parsed.query else ""
    return f"{parsed.path}{query}"


def _tail(location: str) -> str:
    return location.split("?", 1)[0].rstrip("/").split("/")[-1]


def _sha(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _start(client: TestClient, filename: str, payload: bytes) -> tuple[str, str]:
    uploaded = client.post(
        "/api/templates",
        files={"file": (filename, payload, "application/octet-stream")},
        follow_redirects=False,
    )
    assert uploaded.status_code == 303, uploaded.text
    template_id = _tail(_loc(uploaded))
    created = client.post(
        "/api/projects",
        data={"template_id": template_id, "project_name": "web-e2e"},
        follow_redirects=False,
    )
    assert created.status_code == 303, created.text
    return template_id, _tail(_loc(created))


def _generate(client: TestClient, project_id: str, **fields: str):
    data = {"project_title": "웹E2E", "organization_name": "", "user_brief": "", "user_notes": ""}
    data.update(fields)
    return client.post(
        f"/projects/{project_id}/generate",
        data=data,
        follow_redirects=False,
    )


def _route(client: TestClient, project_id: str) -> dict:
    response = client.get(f"/downloads/{project_id}/hwpx_route.json")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["engine"] == "existing hwpx_fill direct-fill"
    return body


def _download_final(client: TestClient, project_id: str, route: dict, dest: Path) -> Path:
    final_name = Path(route["routing"]["final"]).name
    response = client.get(f"/downloads/{project_id}/{final_name}")
    assert response.status_code == 200, response.text
    dest.write_bytes(response.content)
    assert zipfile.is_zipfile(dest)
    with zipfile.ZipFile(dest) as archive:
        assert archive.testzip() is None
    return dest


def _source_unchanged(template_id: str, project_id: str, filename: str, original: bytes) -> None:
    uploaded = main.storage.template_dir(template_id) / filename
    pinned = main.storage.project_dir(project_id) / "template_source.hwpx"
    assert uploaded.read_bytes() == original
    assert pinned.read_bytes() == original
    assert _sha(uploaded.read_bytes()) == _sha(original)


def _no_auto(path: Path) -> None:
    fields = assess_fields(index_hwpx_structure(path))
    assert fields
    assert all(field.decision != "AUTO" and field.auto_write_allowed is False for field in fields)


def _no_staging(project_id: str) -> None:
    output_dir = main.storage.project_dir(project_id) / "output"
    names = [path.name for path in output_dir.iterdir()]
    assert not any(name.startswith(".t02-auth-") for name in names)
    assert not any("__t02_auth__" in name or "__grid_repair__" in name or "__cleanup__" in name for name in names)
    assert not any(name.endswith(".pdf") for name in names)


def _page_lists(client: TestClient, project_id: str, filename: str, *, submittable: bool) -> None:
    page = client.get(f"/projects/{project_id}")
    assert page.status_code == 200
    assert f"/downloads/{project_id}/{filename}" in page.text
    if submittable:
        assert "제출가능: 예" in page.text
    else:
        assert "제출가능: 아니오" in page.text


def test_web_happy_path_returns_intact_filled_package(client, tmp_path):
    src = tmp_path / "form.hwpx"
    _write(src, _section_clean())
    original = src.read_bytes()
    before_header = zipfile.ZipFile(src).read("Contents/header.xml")
    template_id, project_id = _start(client, "form.hwpx", original)
    generated = _generate(client, project_id, organization_name=_FILL, 대표자="침범")
    assert generated.status_code == 303, generated.text
    assert "error=" not in _loc(generated)

    route = _route(client, project_id)
    routing = route["routing"]
    assert routing["ok"] is True
    assert routing["submittable"] is True
    assert routing["final_output_allowed"] is True
    assert Path(routing["final"]).name == "output.hwpx"
    assert RHWP_DISABLED_RENDER_NOTE not in routing["notes"]
    assert routing["native_render"]["render_status"] == "PASS"
    assert routing["native_render"]["reopen_status"] == "PASS"
    assert routing["native_render"]["visual_review"] == "PASS"
    assert routing["native_render"]["severity"] == "PASS"
    assert routing["native_render"]["pixel_reopen_claimed"] is False
    # COM 설치 여부와 rhwp 부재는 별개다. L005 픽셀 PASS는 여기 기대값이 아니다.
    assert route["visual_render"] == (
        "REVIEW_REQUIRED" if hancom_com_available() else "ENVIRONMENT_BLOCKED"
    )
    assert any("대표자" in note and "EXISTING_VALUE" in note for note in routing["notes"])
    assert not any("대표자" in note and note.endswith("GRANT_WRITTEN") for note in routing["notes"])

    final = _download_final(client, project_id, route, tmp_path / "out.hwpx")
    _assert_package(final, _FILL, _EXISTING)
    assert "침범" not in _section_text(final)
    with zipfile.ZipFile(final) as archive:
        assert archive.read("Contents/header.xml") == before_header
    _source_unchanged(template_id, project_id, "form.hwpx", original)
    _no_auto(final)
    _no_staging(project_id)
    _page_lists(client, project_id, "output.hwpx", submittable=True)
    assert not (main.storage.project_dir(project_id) / "output" / "output_DRAFT.hwpx").exists()
    assert not (main.storage.project_dir(project_id) / "output" / "output.docx").exists()

    api = client.post(f"/api/projects/{project_id}/generate")
    assert api.status_code == 200, api.text
    body = api.json()
    assert body["output_hwpx"].endswith("output.hwpx")
    assert Path(body["output_hwpx"]).is_file()
    assert body["output_docx"] == ""
    _source_unchanged(template_id, project_id, "form.hwpx", original)


def test_web_stacked_protections_match_direct_submit(client, tmp_path):
    src = tmp_path / "combined.hwpx"
    _write_combined(src)
    original = src.read_bytes()
    before_cells = _cells(src)
    before_runs = _body_runs(src)
    header = _member(src, "Contents/header.xml")
    package = _member(src, "Contents/content.hpf")
    template_id, project_id = _start(client, "combined.hwpx", original)
    generated = _generate(client, project_id, **_identity())
    assert generated.status_code == 303, generated.text
    assert "error=" not in _loc(generated)

    route = _route(client, project_id)
    routing = route["routing"]
    notes = routing["notes"]
    final = _download_final(client, project_id, route, tmp_path / "filled.hwpx")
    filled = _cells(final)
    _assert_structure_held(before_cells, filled, {
        (0, 1, 1): ["서울"],
        (0, 6, 1): ["홍길동"],
        (0, 7, 1): ["□개인 ■법인"],
        (0, 10, 1): ["침범"],
        (0, 12, 1): ["부산"],
        (0, 13, 1): ["첫번째"],
        (0, 15, 1): ["99"],
        (1, 0, 1): ["자식회사"],
    })
    assert _at(filled, 0, 0, 1)["texts"] == ["기존회사"]
    assert _at(filled, 0, 2, 1)["texts"] == ["000", "억원"]
    assert _at(filled, 0, 3, 1)["texts"] == ["2025년 ", "월 ", "일"]
    assert _at(filled, 0, 4, 1)["texts"] == [""]
    assert _at(filled, 0, 5, 1)["texts"] == [""]
    assert _at(filled, 0, 5, 1)["pics"] == 1
    assert _at(filled, 0, 14, 1)["texts"] == [""]
    assert _at(filled, 3, 0, 1)["texts"] == [""]
    assert _at(filled, 4, 0, 1)["texts"] == [""]
    texts = _all_texts(final)
    for leaked in ("새회사", "김철수", "12", "010", "부모값", "들어가면안됨", "2026년 9월 28일", "홍길동서명"):
        assert leaked not in texts
    assert texts.count("첫번째") == 1
    assert "02" not in texts
    assert _body_runs(final) == before_runs
    assert "[existing] 기업명 row=0 col=1 EXISTING_VALUE" in notes
    assert "[span] 매출 row=2 col=1 UNFILLED" in notes
    assert "[signature] 서명 row=4 col=1 UNFILLED" in notes
    assert "[handwritten] 성명 row=5 col=1 UNFILLED" in notes
    assert "[nested] 겉라벨 row=8 col=1 UNFILLED" in notes
    assert "[nested] 포함표 row=16 col=1 UNFILLED" in notes
    assert "팩스" in routing["residual"]
    assert "직원명" in routing["residual"]
    assert any(note == "[t02] 팩스 AUTHORIZATION_PENDING" for note in notes)
    assert any(note == "[repeated-row] 직원명 row=14 col=1 UNFILLED" for note in notes)
    assert any(note.startswith("[duplicate-label] 팩스 ") for note in notes)
    assert not any("팩스" in note and note.endswith("GRANT_WRITTEN") for note in notes)
    assert "[t02] 속기업명 GRANT_WRITTEN" in notes
    assert "[t02] 기업명 GRANT_WRITTEN" not in notes
    assert not any("서명" in note and note.endswith("GRANT_WRITTEN") for note in notes)
    assert "성명" not in routing["filled"]
    assert "매출" not in routing["filled"]
    assert "겉라벨" not in routing["filled"]
    assert "포함표" not in routing["filled"]
    guides = run_hwpx_acceptance(final)
    assert guides.guides == 1
    assert "작성방법" in guides.guides_samples[0]
    assert routing["submittable"] is False
    assert routing["ok"] is False
    assert Path(routing["final"]).name == "output_DRAFT.hwpx"
    assert client.get(f"/downloads/{project_id}/output.hwpx").status_code == 404
    assert _member(final, "Contents/header.xml") == header
    assert _member(final, "Contents/content.hpf") == package
    assert _foreign_t(final) == 0
    assert RHWP_DISABLED_RENDER_NOTE not in notes
    assert routing["native_render"]["render_status"] == "PASS"
    assert routing["native_render"]["reopen_status"] == "PASS"
    assert routing["native_render"]["visual_review"] == "PASS"
    assert routing["native_render"]["l005_pixel"] != "PASS"
    assert routing["native_render"]["l050_pdf"] != "PASS"
    _source_unchanged(template_id, project_id, "combined.hwpx", original)
    _no_auto(src)
    _no_auto(final)
    _no_staging(project_id)
    _page_lists(client, project_id, "output_DRAFT.hwpx", submittable=False)


def test_web_duplicate_empty_label_writes_nothing(client, tmp_path):
    section = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<hs:sec xmlns:hp="{_HP}" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section">'
        '<hp:p><hp:run><hp:tbl rowCnt="1" colCnt="2"><hp:tr>'
        '<hp:tc><hp:subList><hp:p><hp:run><hp:t>기업명</hp:t></hp:run></hp:p></hp:subList>'
        '<hp:cellAddr rowAddr="0" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="1"/></hp:tc>'
        '<hp:tc><hp:subList><hp:p><hp:run><hp:t></hp:t></hp:run></hp:p></hp:subList>'
        '<hp:cellAddr rowAddr="0" colAddr="1"/><hp:cellSpan rowSpan="1" colSpan="1"/></hp:tc>'
        "</hp:tr></hp:tbl></hp:run></hp:p>"
        '<hp:p><hp:run><hp:tbl rowCnt="1" colCnt="2"><hp:tr>'
        '<hp:tc><hp:subList><hp:p><hp:run><hp:t>기업명</hp:t></hp:run></hp:p></hp:subList>'
        '<hp:cellAddr rowAddr="0" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="1"/></hp:tc>'
        '<hp:tc><hp:subList><hp:p><hp:run><hp:t></hp:t></hp:run></hp:p></hp:subList>'
        '<hp:cellAddr rowAddr="0" colAddr="1"/><hp:cellSpan rowSpan="1" colSpan="1"/></hp:tc>'
        "</hp:tr></hp:tbl></hp:run></hp:p>"
        "</hs:sec>"
    )
    payload = _bare_hwpx(section)
    template_id, project_id = _start(client, "dup.hwpx", payload)
    generated = _generate(client, project_id, organization_name="첫번째")
    assert generated.status_code == 303, generated.text
    assert "error=" not in _loc(generated)
    route = _route(client, project_id)
    notes = route["routing"]["notes"]
    final = _download_final(client, project_id, route, tmp_path / "dup-out.hwpx")
    texts = _all_texts(final)
    assert texts.count("기업명") == 2
    assert "첫번째" not in texts
    assert "기업명" in route["routing"]["residual"]
    assert any(note == "[t02] 기업명 AUTHORIZATION_PENDING" for note in notes)
    assert any(note.startswith("[duplicate-label] 기업명 ") for note in notes)
    assert not any("기업명" in note and note.endswith("GRANT_WRITTEN") for note in notes)
    assert route["routing"]["submittable"] is True, {
        "reason": route["routing"].get("draft_reason"),
        "acceptance": route["routing"].get("acceptance"),
        "status": route["routing"].get("routing_status"),
        "notes": route["routing"].get("notes"),
    }
    _source_unchanged(template_id, project_id, "dup.hwpx", payload)
    _no_auto(final)


def test_web_clean_form_without_render_evidence_is_draft(client, tmp_path, monkeypatch):
    monkeypatch.setattr(native_hwp, "verify_hwpx_native", lambda _path: native_hwp._disabled_render_evidence())
    src = tmp_path / "no-render.hwpx"
    _write(src, _section_clean())
    _template_id, project_id = _start(client, "no-render.hwpx", src.read_bytes())
    generated = _generate(client, project_id, organization_name=_FILL)
    assert generated.status_code == 303, generated.text
    route = _route(client, project_id)
    routing = route["routing"]
    assert routing["ok"] is False
    assert routing["submittable"] is False
    assert Path(routing["final"]).name == "output_DRAFT.hwpx"
    render = next(
        item for item in routing["integrity"]["validators"]
        if item["source_validator"] == "rendering_validator"
    )
    assert render["severity"] == "REVIEW_REQUIRED"
    assert render["defect_code"] == "RENDER_NOT_RUN_REQUIRED"


def test_web_broken_grid_rhwp_absent_keeps_package_and_drafts(client, tmp_path, monkeypatch):
    monkeypatch.setattr(native_hwp, "verify_hwpx_native", lambda _path: native_hwp._disabled_render_evidence())
    src = tmp_path / "broken.hwpx"
    _write(src, _section_protected_broken())
    original = src.read_bytes()
    template_id, project_id = _start(client, "broken.hwpx", original)
    generated = _generate(client, project_id, organization_name=_FILL)
    assert generated.status_code == 303, generated.text
    assert "error=" not in _loc(generated)
    route = _route(client, project_id)
    routing = route["routing"]
    assert routing["ok"] is False
    assert routing["submittable"] is False
    assert routing["repair"] == {}
    assert RHWP_DISABLED_REPAIR_NOTE in routing["notes"]
    assert RHWP_DISABLED_RENDER_NOTE in routing["notes"]
    assert routing["native_render"]["reopen_status"] == "NOT_RUN"
    assert routing["native_render"]["render_status"] == "NOT_RUN"
    assert routing["native_render"]["l005_pixel"] == "NOT_RUN"
    assert routing["native_render"]["severity"] == "PASS"
    assert Path(routing["final"]).name == "output_DRAFT.hwpx"
    assert client.get(f"/downloads/{project_id}/output.hwpx").status_code == 404
    final = _download_final(client, project_id, route, tmp_path / "broken-out.hwpx")
    section = _assert_package(final, _FILL, _SIGNATURE, _CHECKBOX, _EXISTING, _DATE)
    assert 'rowAddr="2"' not in section
    _source_unchanged(template_id, project_id, "broken.hwpx", original)
    _no_staging(project_id)
    _page_lists(client, project_id, "output_DRAFT.hwpx", submittable=False)


def test_web_colored_template_keeps_color_and_submit_name(client, tmp_path):
    """양식 유색(FF0000)은 지우지 않고, 그 색만으로 _DRAFT 가 되지 않는다."""
    src = tmp_path / "colored.hwpx"
    _write(src, _section_clean(), colored=True)
    original = src.read_bytes()
    template_id, project_id = _start(client, "colored.hwpx", original)
    generated = _generate(client, project_id, organization_name=_FILL, 대표자="침범")
    assert generated.status_code == 303, generated.text
    assert "error=" not in _loc(generated)
    route = _route(client, project_id)
    routing = route["routing"]
    assert routing["ok"] is True
    assert routing["submittable"] is True
    assert Path(routing["final"]).name == "output.hwpx"
    assert client.get(f"/downloads/{project_id}/output_DRAFT.hwpx").status_code == 404
    final = _download_final(client, project_id, route, tmp_path / "colored-out.hwpx")
    _assert_package(final, _FILL, _EXISTING)
    assert b"FF0000" in zipfile.ZipFile(final).read("Contents/header.xml")
    assert "침범" not in _section_text(final)
    assert routing["native_render"]["l005_pixel"] != "PASS"
    _source_unchanged(template_id, project_id, "colored.hwpx", original)


def test_web_bad_zip_is_400_and_leaves_no_output(client):
    payload = b"this is not a zip"
    template_id, project_id = _start(client, "form.hwpx", payload)
    api = client.post(f"/api/projects/{project_id}/generate")
    assert api.status_code == 400, api.text
    assert "HWPX" in api.json()["detail"] or "ZIP" in api.json()["detail"]
    assert client.get(f"/downloads/{project_id}/output.hwpx").status_code == 404
    assert client.get(f"/downloads/{project_id}/output_DRAFT.hwpx").status_code == 404
    output_dir = main.storage.project_dir(project_id) / "output"
    assert list(output_dir.glob("*.hwpx")) == []
    _source_unchanged(template_id, project_id, "form.hwpx", payload)

    form = _generate(client, project_id, organization_name=_FILL)
    assert form.status_code == 303, form.text
    location = _loc(form)
    assert "error=" in location
    page = client.get(location)
    assert page.status_code == 200
    assert "HWPX" in page.text or "ZIP" in page.text
    assert list(output_dir.glob("*.hwpx")) == []
    _source_unchanged(template_id, project_id, "form.hwpx", payload)


def test_web_upload_refuses_announcement_hwpx(client):
    response = client.post(
        "/api/templates",
        files={"file": ("모집공고.hwpx", b"PK", "application/octet-stream")},
        follow_redirects=False,
    )
    assert response.status_code == 400, response.text
    assert "공고" in response.json()["detail"]


def test_api_generate_maps_missing_file_to_400(client, monkeypatch):
    def missing(_project_id: str):
        raise FileNotFoundError("입력 파일이 없습니다: form.hwpx")

    monkeypatch.setattr(main.project_service, "generate", missing)
    response = client.post("/api/projects/prj_missing/generate")
    assert response.status_code == 400, response.text
    body = response.json()
    assert body["detail"]
    assert "없습니다" in body["detail"]


def test_operator_hwpx_keeps_existing_blocks_signature_and_synonym(client, tmp_path):
    section = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        f'<hs:sec xmlns:hp="{_HP}" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section">'
        + _pair_table("기업명", "기존회사")
        + _pair_table("상호", "")
        + _pair_table("주소", "")
        + _pair_table("서명", "")
        + "</hs:sec>"
    )
    payload = _bare_hwpx(section)
    facts = _facts_docx()
    response = client.post(
        "/console/documents/write",
        data={"project_title": "운영콘솔", "organization_name": "새회사", "instruction": ""},
        files=[
            ("template_file", ("form.hwpx", payload, "application/octet-stream")),
            ("reference_files", ("facts.docx", facts, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ],
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text
    location = _loc(response)
    assert "error=" not in location, location
    project_id = _tail(location)
    page = client.get(location)
    assert page.status_code == 200, page.text
    route = _route(client, project_id)
    final = _download_final(client, project_id, route, tmp_path / "op.hwpx")
    text = _section_text(final)
    assert "기존회사" in text
    assert "서울특별시" in text
    assert "새회사" not in text
    assert "홍길동서명" not in text
    assert any("서명" in note and "UNFILLED" in note for note in route["routing"]["notes"])
    assert any("기업명" in note and "EXISTING_VALUE" in note for note in route["routing"]["notes"])
    pinned = main.storage.project_dir(project_id) / "template_source.hwpx"
    assert pinned.read_bytes() == payload
    assert not (main.storage.project_dir(project_id) / "output" / "output.docx").exists()
    _no_auto(final)


def _section_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("Contents/section")
        )


def _pair_table(label: str, value: str) -> str:
    return (
        '<hp:p><hp:run><hp:tbl rowCnt="1" colCnt="2"><hp:tr>'
        '<hp:tc><hp:subList><hp:p><hp:run><hp:t>' + label + "</hp:t></hp:run></hp:p></hp:subList>"
        '<hp:cellAddr rowAddr="0" colAddr="0"/><hp:cellSpan rowSpan="1" colSpan="1"/></hp:tc>'
        '<hp:tc><hp:subList><hp:p><hp:run><hp:t>' + value + "</hp:t></hp:run></hp:p></hp:subList>"
        '<hp:cellAddr rowAddr="0" colAddr="1"/><hp:cellSpan rowSpan="1" colSpan="1"/></hp:tc>'
        "</hp:tr></hp:tbl></hp:run></hp:p>"
    )


def _bare_hwpx(section: str) -> bytes:
    header = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head"><hh:refList>'
        '<hh:charProperties itemCnt="1"><hh:charPr id="0" textColor="000000"/>'
        "</hh:charProperties></hh:refList></hh:head>"
    )
    hpf = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<opf:package xmlns:opf="http://www.idpf.org/2007/opf">'
        '<opf:manifest><opf:item id="s0" href="Contents/section0.xml" media-type="application/xml"/>'
        '</opf:manifest><opf:spine><opf:itemref idref="s0"/></opf:spine></opf:package>'
    )
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        info = zipfile.ZipInfo("mimetype")
        info.compress_type = zipfile.ZIP_STORED
        archive.writestr(info, b"application/hwp+zip")
        archive.writestr("Contents/header.xml", header.encode("utf-8"))
        archive.writestr("Contents/section0.xml", section.encode("utf-8"))
        archive.writestr("Contents/content.hpf", hpf.encode("utf-8"))
    return buffer.getvalue()


def _facts_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    rows = (("기업명", "새회사"), ("주소", "서울특별시"), ("서명", "홍길동서명"))
    for index, (label, value) in enumerate(rows):
        table.cell(index, 0).text = label
        table.cell(index, 1).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
