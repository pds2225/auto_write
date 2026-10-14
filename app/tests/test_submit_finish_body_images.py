"""F3 사업계획서 본문 자동 채움 + F4 섹션별 이미지 삽입 (가짜 데이터만 사용)."""
from __future__ import annotations

import json
import os
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from PIL import Image

from auto_write.services.company_identity import company_profile_identity
from auto_write.services.hwpx_form_diff import compare_hwpx_forms
from auto_write.services.hwpx_narrative_source import direct_cell_plan
from auto_write.services.hwpx_pic_insert import insert_cell_reference_images
from auto_write.services.submit_finish.body_fill import heading_key, load_json_reference
from auto_write.services.submit_finish.section_images import plan_section_images
from core.docx.services.hwpx_analysis_adapter import index_hwpx_structure
from core.docx.services.hwpx_fill import commit_guidance_value_writes
from test_hwpx_engine_coverage_synthetic import LABELS, NS, HP, narrative_table, package, root, table
from test_hwpx_round2_fill import _cell, _row, _run

# 실제 사업계획서처럼 번호·기호가 붙은 제목(마포 재현 입력과 같은 모양).
NUMBERED = ['1. 아이디어 제안배경', '2. 신청자 역량', '3. 창업 아이디어',
            '가. 핵심내용', '나) 우수성 및 차별성', '(다) 시장 경쟁력', '4. 기대효과']
MASTER = {'company_key': '테스트회사', 'fields': {
    '기업명': {'value': '테스트회사', 'confidence': 'medium', 'confirmed': False, 'sources': []},
    '대표자': {'value': '홍길동', 'confidence': 'high', 'confirmed': False, 'sources': []},
    '주소': {'value': 'A', 'confidence': 'conflict', 'confirmed': False, 'sources': []}},
    'conflicts': [{'field': '주소'}], 'missing': [], 'sources': [], 'updated_at': ''}


def _bizplan(path: Path, heading_style: bool) -> Path:
    doc = Document()
    for heading in NUMBERED:
        if heading_style:
            doc.add_heading(heading, level=1)
        else:
            doc.add_paragraph(heading)
        doc.add_paragraph(f'{heading_key(heading)} 가짜 본문 1')
        doc.add_paragraph(f'{heading_key(heading)} 가짜 본문 2')
    doc.save(path)
    return path


def _png(path: Path, size, fmt='PNG') -> Path:
    Image.new('RGB', size, 'red').save(path, format=fmt)
    return path


def _plan(src: Path, refs=(), identity=None, meta=None):
    index = index_hwpx_structure(src)
    return direct_cell_plan(index, identity or {}, list(refs), meta or {})


def _commit(tmp_path: Path, src: Path, grants, values) -> Path:
    out = tmp_path / 'filled.hwpx'
    commit_guidance_value_writes(src, out, grants, values)
    return out


# ---------------- F3 ----------------

@pytest.mark.parametrize('heading_style', [True, False])
def test_numbered_bizplan_headings_fill_every_narrative(tmp_path, heading_style):
    """수정 전: 7칸 전부 NO_SOURCE. 수정 후: 필수 본문 누락 0·안내문 0·라벨칸 보존."""
    src = package(tmp_path, narrative_table())
    original = src.read_bytes()
    ref = _bizplan(tmp_path / 'plan.docx', heading_style)
    grants, values, notes, _ = _plan(src, [ref])
    assert not [n for n in notes if 'NO_SOURCE' in n], notes
    out = _commit(tmp_path, src, grants, values)
    text = ''.join(root(out).itertext())
    for label in LABELS:
        assert f'{heading_key(label)} 가짜 본문 1' in text, label
    assert '에 대해 작성' not in text
    old, new = root(src), root(out)
    for a, b in zip(old.xpath('//hp:tc[position()=1]', namespaces=NS), new.xpath('//hp:tc[position()=1]', namespaces=NS)):
        assert etree.tostring(a) == etree.tostring(b)  # 라벨칸(기존 내용) 손실 0
    assert compare_hwpx_forms(src, out).form_intact
    assert src.read_bytes() == original


def test_heading_key_strips_only_enumeration_prefix():
    assert heading_key('1. 아이디어 제안배경') == heading_key('아이디어 제안배경')
    assert heading_key('□ 핵심내용') == heading_key('(가) 핵심내용') == heading_key('핵심내용')
    assert heading_key('3D 프린팅') != heading_key('D 프린팅')  # 숫자 시작 단어는 보존
    assert heading_key('아이디어 제안배경 및 필요성') != heading_key('아이디어 제안배경')  # 정확 일치 유지


def test_company_master_json_and_bizplan_data_reference(tmp_path):
    """회사 마스터 JSON + 사업계획 데이터 JSON(가짜) → 본문·반복표 채움, 사전 원문/충돌값 미기입."""
    headers = ['권리구분', '등록(출원)명칭', '등록(출원)번호', '등록(출원)인', '기타']
    label = _cell(0, 0, _run('지식재산권')).replace('rowSpan="1"', 'rowSpan="3"')
    rows = ['<hp:tr>' + label + ''.join(_cell(i + 1, 0, _run(h)) for i, h in enumerate(headers)) + '</hp:tr>']
    rows += ['<hp:tr>' + ''.join(_cell(i + 1, r, _run('')) for i in range(5)) + '</hp:tr>' for r in (1, 2)]
    src = package(tmp_path, narrative_table() + table(rows))
    records = [dict(zip(headers, ['특허', f'테스트 시스템 {n}', f'10-2026-000000{n}', '홍길동', '출원'])) for n in (1, 2)]
    payload = dict(MASTER, business_plan={label: f'{label} JSON 본문' for label in LABELS}, hwpx_records={'지식재산권': records})
    ref = tmp_path / 'company_master.json'
    ref.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    loaded = load_json_reference(ref)
    assert loaded['identity']['대표자'] == '홍길동' and '주소' not in loaded['identity']
    grants, values, notes, _ = _plan(src, [ref])
    assert not [n for n in notes if 'NO_SOURCE' in n or 'MISMATCH' in n or 'OVERFLOW' in n], notes
    out = _commit(tmp_path, src, grants, values)
    text = ''.join(root(out).itertext())
    for label in LABELS:
        assert f'{label} JSON 본문' in text
    # 반복 데이터 수(2) = 채운 표 행 수(2)
    assert text.count('특허') == 2 and '10-2026-0000001' in text and '10-2026-0000002' in text
    assert "{'value'" not in text and 'confidence' not in text
    assert compare_hwpx_forms(src, out).form_intact


def test_master_dict_values_never_stringified():
    identity = company_profile_identity(MASTER['fields'])
    assert identity == {'기업명': '테스트회사', '대표자': '홍길동'}  # 충돌 주소는 비움


def test_existing_user_value_kept_and_missing_data_reported(tmp_path):
    body = table([_row(0, '핵심내용', '사용자가 이미 쓴 내용'), _row(1, '시장 경쟁력', '')])
    src = package(tmp_path, body)
    ref = _bizplan(tmp_path / 'plan.docx', True)
    grants, values, notes, _ = _plan(src, [ref])
    assert [g.field_label for g in grants] == ['시장 경쟁력']  # 기존 값 칸은 승인조차 안 됨
    out = _commit(tmp_path, src, grants, values)
    assert '사용자가 이미 쓴 내용' in ''.join(root(out).itertext())
    grants, values, notes, _ = _plan(src, [])  # 데이터 없음 → 날조 없이 비움 + 사유
    assert not grants and any('시장 경쟁력 NO_SOURCE' in n for n in notes)


# ---------------- F4 ----------------

def _section_setup(tmp_path):
    src = package(tmp_path, narrative_table())
    ref = _bizplan(tmp_path / 'plan.docx', True)
    grants, values, _, docx_images = _plan(src, [ref])
    assert not docx_images
    return src, _commit(tmp_path, src, grants, values), grants, values


def _ratio_ok(pic, px):
    sz = pic.find(f'{{{HP}}}sz')
    w, h = int(sz.get('width')), int(sz.get('height'))
    assert w > 0 and h > 0
    assert abs(w / h - px[0] / px[1]) < 0.01, (w, h, px)  # 가로·세로 모두 원본 비율(L142)
    return w


def test_section_images_inserted_per_section_ratio_and_idempotent(tmp_path):
    src, filled, grants, values = _section_setup(tmp_path)
    wide = _png(tmp_path / 'wide.png', (600, 300))
    tall = _png(tmp_path / 'tall.jpg', (300, 900), 'JPEG')
    entries = [{'anchor': '3. 창업 아이디어', 'image': str(wide), 'caption': '그림 1. 서비스 구조'},
               {'anchor': '기대효과', 'image': 'tall.jpg', 'caption': '그림 2. 기대효과'},
               {'anchor': '창업 아이디어', 'image': str(wide)}]  # 같은 칸·같은 그림 반복
    images, notes = plan_section_images(entries, grants, values, current_hwpx=filled, base_dirs=[tmp_path])
    assert sum(len(v) for v in images.values()) == 2
    assert any('DUPLICATE_SKIPPED' in n for n in notes)
    final = tmp_path / 'final.hwpx'
    insert_cell_reference_images(filled, final, images, authorization_source=src, values=values)
    pics = root(final).xpath('//hp:pic', namespaces=NS)
    assert len(pics) == 2
    widths = [_ratio_ok(pics[0], (600, 300)), _ratio_ok(pics[1], (300, 900))]
    assert all(w <= 10000 for w in widths)  # 칸 폭(=쪽 영역 안) 상한
    text = ''.join(root(final).itertext())
    assert '그림 1. 서비스 구조' in text and '그림 2. 기대효과' in text
    with ZipFile(final) as z:
        assert len([n for n in z.namelist() if n.startswith('BinData/')]) == 2
        manifest = etree.fromstring(z.read('Contents/content.hpf'))
        assert len(manifest.xpath('//*[local-name()="item" and @isEmbeded="1"]')) == 2
    assert compare_hwpx_forms(src, final).form_intact
    # 같은 입력을 결과물에 다시 적용 → 새로 넣을 그림 0(멱등)
    again, notes = plan_section_images(entries, grants, values, current_hwpx=final, base_dirs=[tmp_path])
    assert not any(again.values())
    assert sum('DUPLICATE_SKIPPED' in n for n in notes) == 3


def test_section_images_missing_or_unlocatable_are_review(tmp_path):
    src, filled, grants, values = _section_setup(tmp_path)
    img = _png(tmp_path / 'ok.png', (100, 100))
    gif = tmp_path / 'bad.gif'
    Image.new('RGB', (10, 10)).save(gif, format='GIF')
    entries = [{'anchor': '핵심내용', 'image': str(tmp_path / 'none.png')},
               {'anchor': '없는 섹션', 'image': str(img)},
               {'anchor': '우수성 및 차별성', 'image': str(gif)}]
    images, notes = plan_section_images(entries, grants, values, current_hwpx=filled)
    assert not any(images.values())
    for code in ('IMAGE_FILE_MISSING', 'ANCHOR_NOT_FOUND', 'IMAGE_UNSUPPORTED'):
        assert any(code in n and n.endswith('REVIEW') for n in notes), (code, notes)
    # 본문이 안 채워진 칸(데이터 없음)은 이미지도 넣지 않는다
    images, notes = plan_section_images([{'anchor': '핵심내용', 'image': str(img)}], grants, {g: '' for g in grants})
    assert not images and 'ANCHOR_NOT_WRITTEN' in notes[0]


from test_hwpx_web_e2e_gate import web, client, _route, _tail, _loc  # noqa: E402,F401


def _no_com(monkeypatch):
    """한글 COM 을 띄우지 않는다(사용자 Hwp 프로세스 보호)."""
    import core.docx.services.hwp_docx_convert as core_com
    monkeypatch.setattr(core_com, 'hancom_com_available', lambda: False)
    try:
        import auto_write.services.hwp_docx_convert as shim
        monkeypatch.setattr(shim, 'hancom_com_available', lambda: False, raising=False)
    except Exception:
        pass


def _console(client, form_bytes, refs):
    files = [('template_file', ('form.hwpx', form_bytes, 'application/octet-stream'))]
    files += [('reference_files', (name, data, 'application/octet-stream')) for name, data in refs]
    response = client.post('/console/documents/write', data={'project_title': 'F3F4 합성', 'organization_name': '', 'instruction': ''}, files=files, follow_redirects=False)
    assert response.status_code == 303, response.text
    location = _loc(response)
    assert 'error=' not in location, location
    return _route(client, _tail(location))


def _section_json(tmp_path, anchors):
    img = _png(tmp_path / 'fig.png', (800, 400))
    payload = dict(MASTER, section_images=[{'anchor': a, 'image': 'fig.png', 'caption': f'{a} 그림'} for a in anchors])
    return [('company_master.json', json.dumps(payload, ensure_ascii=False).encode('utf-8')), ('fig.png', img.read_bytes())]


def test_console_synthetic_form_numbered_plan_and_section_image(client, monkeypatch, tmp_path):
    """콘솔 경로 실사용: 번호 제목 사업계획서 + 마스터 JSON + 섹션 이미지(opt-in)."""
    _no_com(monkeypatch)
    src = package(tmp_path, narrative_table())
    plan = _bizplan(tmp_path / 'plan.docx', False)
    refs = [('plan.docx', plan.read_bytes())] + _section_json(tmp_path, ['핵심내용', '없는 섹션'])
    route = _console(client, src.read_bytes(), refs)
    out = Path(route['routing']['final'])
    text = ''.join(root(out).itertext())
    for label in LABELS:
        assert f'{heading_key(label)} 가짜 본문 1' in text, label
    assert len(root(out).xpath('//hp:pic', namespaces=NS)) == 1
    notes = route['routing']['notes']
    assert any('없는 섹션 ANCHOR_NOT_FOUND REVIEW' in n for n in notes)
    assert any('ANCHOR_NOT_FOUND' in r for r in route['routing']['residual'])
    # 설정 없음 → 기존 동작(그림 없음)
    route = _console(client, src.read_bytes(), [('plan.docx', plan.read_bytes())])
    assert not root(Path(route['routing']['final'])).xpath('//hp:pic', namespaces=NS)


def test_real_mapo_form_body_and_section_images(client, monkeypatch, tmp_path):
    """마포 실양식 사본(AUTO_WRITE_MAPO_FORM)에서 가짜 데이터로 본문 7칸 + 섹션 이미지."""
    source = Path(os.environ.get('AUTO_WRITE_MAPO_FORM', ''))
    if not source.is_file():
        pytest.skip('AUTO_WRITE_MAPO_FORM 없음')
    _no_com(monkeypatch)
    before = source.read_bytes()
    plan = _bizplan(tmp_path / 'plan.docx', False)
    refs = [('plan.docx', plan.read_bytes())] + _section_json(tmp_path, ['3. 창업 아이디어', '기대효과'])
    route = _console(client, before, refs)
    out = Path(route['routing']['final'])
    text = ''.join(root(out).itertext())
    for label in LABELS:
        assert f'{heading_key(label)} 가짜 본문 1' in text, label
    for phrase in ('아이디어 착안 배경, 아이디어를 통해', '사용자나 사회에 미치는 긍정적 변화'):
        assert phrase not in text  # 값 넣은 칸의 안내문 제거
    with ZipFile(source) as z:
        original_pics = len(etree.fromstring(z.read('Contents/section0.xml')).xpath('//hp:pic', namespaces=NS))
    assert len(root(out).xpath('//hp:pic', namespaces=NS)) == original_pics + 2
    compare = compare_hwpx_forms(source, out)
    assert compare.form_intact, compare.as_dict()
    assert source.read_bytes() == before
