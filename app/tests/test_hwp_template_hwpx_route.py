"""HWP 양식 업로드는 HWPX direct-fill로 고정한다. 변환기가 없으면 DOCX로 내려가지 않는다."""

from __future__ import annotations

import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urlparse

from docx import Document
from fastapi.testclient import TestClient

import auto_write.main as main
import auto_write.operator_main as operator_main
from auto_write.config import ensure_directories
from auto_write.models import TemplateProfile
from auto_write.services.evidence_service import EvidenceService
from auto_write.services.hwp_docx_convert import ConvertReport, HWP_TO_HWPX_UNAVAILABLE_NOTE
from auto_write.services.image_service import ImageService
from auto_write.services.openai_client import OpenAIService
from auto_write.services.project_service import ProjectService
from auto_write.services.qa_service import QAService
from auto_write.services.render_service import RenderService
from auto_write.storage import Storage
from test_project_service_safety import _minimal_hwpx_bytes, build_settings


def _service(tmp_path: Path) -> tuple[ProjectService, Storage]:
    root = tmp_path / "appdata"
    settings = build_settings(root)
    ensure_directories(settings)
    storage = Storage(settings)
    openai_service = OpenAIService(settings)
    service = ProjectService(
        storage=storage,
        openai_service=openai_service,
        evidence_service=EvidenceService(openai_service),
        image_service=ImageService(openai_service),
        render_service=RenderService(),
        qa_service=QAService(),
    )
    return service, storage


def _ok_report(src: Path, dst: Path) -> ConvertReport:
    return ConvertReport(
        direction="hwp->hwpx",
        method="hancom_com",
        ok=True,
        output=str(dst),
        notes=["한글 COM으로 HWPX를 저장했습니다. 원본 HWP는 수정하지 않았습니다."],
    )


def _fake_converter(src, dst=None):
    src_path = Path(src)
    dst_path = Path(dst) if dst else src_path.with_suffix(".hwpx")
    original = src_path.read_bytes()
    dst_path.write_bytes(_minimal_hwpx_bytes())
    assert src_path.read_bytes() == original
    return _ok_report(src_path, dst_path)


def _fail_converter(src, dst=None):
    return ConvertReport(
        direction="hwp->hwpx",
        method="",
        ok=False,
        output=str(dst or ""),
        notes=[HWP_TO_HWPX_UNAVAILABLE_NOTE],
    )


def test_hwp_upload_with_converter_pins_source_hwpx_and_generates_hwpx(tmp_path, monkeypatch):
    service, storage = _service(tmp_path)
    original = b"fake-hwp-template"
    monkeypatch.setattr(
        "auto_write.services.hwp_docx_convert.hwp_to_hwpx",
        _fake_converter,
    )
    monkeypatch.setattr(
        "auto_write.services.project_service.ensure_template_docx",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("DOCX 변환 금지")),
    )

    profile = service.analyze_uploaded_template("form.hwp", original)
    project_id = service.create_project(profile.template_id, "HWP 직접 작성")
    service.save_project_form(
        project_id=project_id,
        answers={},
        project_title="HWP 직접 작성",
        organization_name="직접입력(주)",
        evidence_topics="",
        reference_files=[],
    )
    artifacts = service.generate(project_id)

    uploaded = storage.template_dir(profile.template_id) / "form.hwp"
    assert uploaded.read_bytes() == original
    assert profile.template_name == "form.hwp"
    assert profile.source_docx == ""
    assert profile.source_hwpx.endswith("form.hwpx")
    assert Path(profile.source_hwpx).read_bytes() == _minimal_hwpx_bytes()
    assert profile.native_source["conversion"] == "hancom_com"
    assert artifacts.output_docx == ""
    assert Path(artifacts.output_hwpx).is_file()
    assert Path(artifacts.output_hwpx).suffix.lower() == ".hwpx"
    assert Path(artifacts.results_hwpx).is_file()
    output_dir = storage.project_dir(project_id) / "output"
    assert (output_dir / "output.hwpx").is_file()
    assert not (output_dir / "output.docx").exists()
    assert not list(storage.template_dir(profile.template_id).glob("*.docx"))
    route = (output_dir / "hwpx_route.json").read_text(encoding="utf-8")
    assert "existing hwpx_fill direct-fill" in route
    pinned = storage.project_dir(project_id) / "template_source.hwpx"
    assert pinned.read_bytes() == _minimal_hwpx_bytes()
    assert uploaded.read_bytes() == original


def test_hwp_upload_without_converter_does_not_fall_back_to_docx(tmp_path, monkeypatch):
    service, storage = _service(tmp_path)
    original = b"fake-hwp-template"
    monkeypatch.setattr(
        "auto_write.services.hwp_docx_convert.hwp_to_hwpx",
        _fail_converter,
    )

    def _docx_path(*_args, **_kwargs):
        raise AssertionError("변환 불가 HWP가 DOCX 분석으로 내려가면 안 됩니다.")

    monkeypatch.setattr("auto_write.services.project_service.ensure_template_docx", _docx_path)

    try:
        service.analyze_uploaded_template("form.hwp", original)
    except ValueError as exc:
        message = str(exc)
    else:
        raise AssertionError("HWPX 변환 실패는 오류로 끝나야 합니다.")

    assert "HWPX" in message
    assert "DOCX로 진행" in message
    assert "DOCX로 진행하지 않았습니다" in message
    templates = list(storage.settings.template_root.glob("*"))
    assert templates
    for folder in templates:
        assert (folder / "form.hwp").read_bytes() == original
        assert list(folder.glob("*.docx")) == []
        assert list(folder.glob("*.hwpx")) == []
        assert not (folder / "template_profile.json").exists()


def test_generate_refuses_hwp_profile_without_source_hwpx(tmp_path, monkeypatch):
    service, storage = _service(tmp_path)
    template_id = "tpl_hwp_plain"
    storage.template_dir(template_id).mkdir(parents=True)
    storage.save_template_profile(
        TemplateProfile(
            template_id=template_id,
            template_name="form.hwp",
            source_docx="",
            source_hwpx="",
        )
    )
    project_id = service.create_project(template_id, "변환 없는 HWP")

    def _render(*_args, **_kwargs):
        raise AssertionError("HWP 프로필이 DOCX 렌더로 들어가면 안 됩니다.")

    monkeypatch.setattr(service, "_render_and_publish", _render)
    monkeypatch.setattr(
        service,
        "_generate_hwpx_direct",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("source_hwpx 없이 direct-fill 금지")),
    )

    try:
        service.generate(project_id)
    except ValueError as exc:
        message = str(exc)
    else:
        raise AssertionError("source_hwpx 없는 HWP 생성은 막아야 합니다.")

    assert "HWPX" in message
    assert "DOCX로 진행" in message
    assert list((storage.project_dir(project_id) / "output").glob("*.docx")) == []


def test_docx_template_upload_unchanged(tmp_path, monkeypatch):
    service, _storage = _service(tmp_path)
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("테스트 사업계획서")
    document.save(buffer)
    monkeypatch.setattr(
        "auto_write.services.hwp_docx_convert.hwp_to_hwpx",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("DOCX 업로드가 HWP 변환을 타면 안 됩니다.")),
    )

    profile = service.analyze_uploaded_template("정부 지원서 2026.docx", buffer.getvalue())

    assert profile.template_name == "정부 지원서 2026.docx"
    assert profile.source_hwpx == ""
    assert profile.source_docx.endswith("정부 지원서 2026.docx")
    assert Path(profile.source_docx).is_file()
    assert not profile.native_source.get("hwp_docx_opt_in")


def test_hwp_docx_opt_in_stays_on_docx_route(tmp_path, monkeypatch):
    service, storage = _service(tmp_path)
    original = b"fake-hwp-template"

    def _ensure(path: Path):
        converted = path.with_name(f"{path.stem}_converted.docx")
        document = Document()
        document.add_paragraph("명시적 DOCX 진행")
        document.save(converted)
        return converted, ["사용자가 DOCX 진행을 선택했습니다."]

    monkeypatch.setattr("auto_write.services.project_service.ensure_template_docx", _ensure)
    monkeypatch.setattr(
        "auto_write.services.hwp_docx_convert.hwp_to_hwpx",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("명시적 DOCX 진행은 HWPX 변환을 타지 않습니다.")),
    )

    profile = service.analyze_uploaded_template(
        "form.hwp",
        original,
        allow_docx_for_hwp=True,
    )
    project_id = service.create_project(profile.template_id, "DOCX 명시")
    loaded = service.load_profile_for_project(project_id)
    project_input = storage.load_project_input(project_id)

    assert (storage.template_dir(profile.template_id) / "form.hwp").read_bytes() == original
    assert profile.source_hwpx == ""
    assert profile.source_docx.endswith("_converted.docx")
    assert profile.native_source.get("hwp_docx_opt_in") is True
    assert loaded.native_source.get("hwp_docx_opt_in") is True
    sentinel = object()
    monkeypatch.setattr(service, "_render_and_publish", lambda *_args, **_kwargs: sentinel)
    monkeypatch.setattr(
        service,
        "_generate_hwpx_direct",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("opt-in은 direct-fill이 아닙니다.")),
    )
    assert service._hangul_direct_or_refuse(project_id, loaded, project_input) is None
    assert service.generate(project_id) is sentinel


def _isolate_web(monkeypatch, tmp_path: Path):
    root = tmp_path / "workspace"
    settings = replace(
        main.storage.settings,
        workspace_root=root,
        template_root=root / "templates",
        project_root=root / "projects",
        results_root=tmp_path / "results",
    )
    for path in (settings.template_root, settings.project_root, settings.results_root):
        path.mkdir(parents=True)
    monkeypatch.setattr(main.storage, "settings", settings)
    return settings


def test_console_hwp_upload_offers_hwpx_download(tmp_path, monkeypatch):
    _isolate_web(monkeypatch, tmp_path)
    monkeypatch.setattr("auto_write.services.hwp_docx_convert.hwp_to_hwpx", _fake_converter)
    client = TestClient(main.app, raise_server_exceptions=False)
    home = client.get("/console")
    assert home.status_code == 200
    assert "DOCX로 진행" in home.text
    assert 'name="allow_docx_for_hwp"' in home.text

    response = client.post(
        "/console/documents/write",
        data={"project_title": "한글양식", "organization_name": "직접입력(주)", "instruction": "사실만"},
        files={"template_file": ("form.hwp", b"fake-hwp-template", "application/octet-stream")},
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text
    location = response.headers["location"]
    assert "error=" not in location, unquote(location)
    project_id = urlparse(location).path.rstrip("/").split("/")[-1]
    page = client.get(location)
    assert page.status_code == 200
    assert f"/console/results/{project_id}/download/output.hwpx" in page.text
    assert "HWPX 다운로드" in page.text
    assert f"/console/results/{project_id}/download/output.docx" not in page.text

    output_dir = main.storage.project_dir(project_id) / "output"
    (output_dir / "output.docx").write_bytes(b"not-the-hangul-result")
    again = client.get(f"/console/results/{project_id}")
    assert f"/console/results/{project_id}/download/output.hwpx" in again.text
    assert f"/console/results/{project_id}/download/output.docx" not in again.text
    downloaded = client.get(f"/console/results/{project_id}/download/output.hwpx")
    assert downloaded.status_code == 200
    assert downloaded.content == (output_dir / "output.hwpx").read_bytes()
    assert downloaded.content != b"not-the-hangul-result"
    uploaded = list(main.storage.settings.template_root.glob("*/form.hwp"))
    assert uploaded and uploaded[0].read_bytes() == b"fake-hwp-template"


def test_console_hwp_without_converter_shows_korean_error(tmp_path, monkeypatch):
    _isolate_web(monkeypatch, tmp_path)
    monkeypatch.setattr("auto_write.services.hwp_docx_convert.hwp_to_hwpx", _fail_converter)
    client = TestClient(main.app, raise_server_exceptions=False)
    response = client.post(
        "/console/documents/write",
        data={"project_title": "한글양식", "organization_name": "직접입력(주)", "instruction": "사실만"},
        files={"template_file": ("form.hwp", b"fake-hwp-template", "application/octet-stream")},
        follow_redirects=False,
    )
    assert response.status_code == 303
    location = response.headers["location"]
    assert location.startswith("/console?error=")
    message = unquote(urlparse(location).query.split("=", 1)[1])
    assert "HWPX" in message
    assert "DOCX로 진행" in message
    assert "DOCX로 진행하지 않았습니다" in message
    page = client.get(location)
    assert page.status_code == 200
    assert "HWPX" in page.text
    assert "DOCX로 진행" in page.text
    assert list(main.storage.settings.template_root.glob("*/*.docx")) == []
    assert list(main.storage.settings.project_root.glob("*/output/output.docx")) == []
    assert operator_main._form_checked("1") is True
    assert operator_main._form_checked("") is False


def test_console_docx_download_stays_docx_when_hwpx_also_exists(tmp_path, monkeypatch):
    _isolate_web(monkeypatch, tmp_path)
    project_id = "prj_docx_keep"
    project = main.storage.project_dir(project_id)
    output = project / "output"
    output.mkdir(parents=True)
    document = Document()
    document.add_paragraph("워드 양식 결과")
    document.save(output / "output.docx")
    (output / "output.hwpx").write_bytes(_minimal_hwpx_bytes())
    snapshot = {
        "template_id": "tpl_docx",
        "template_name": "form.docx",
        "source_docx": "form.docx",
        "source_hwpx": "",
        "native_source": {},
    }
    (project / "template_snapshot.json").write_text(
        json.dumps(snapshot),
        encoding="utf-8",
    )
    client = TestClient(main.app, raise_server_exceptions=False)
    page = client.get(f"/console/results/{project_id}")
    assert page.status_code == 200
    assert f"/console/results/{project_id}/download/output.docx" in page.text
    assert "DOCX 다운로드" in page.text
    assert f"/console/results/{project_id}/download/output.hwpx" not in page.text
